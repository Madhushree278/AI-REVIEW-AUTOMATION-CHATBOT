import os
import zipfile
from typing import Tuple, Dict, Any, List
from pathlib import Path
import pypdf
import docx

CODE_EXTENSIONS = {
    ".py", ".js", ".ts", ".jsx", ".tsx", ".html", ".css", ".java",
    ".cpp", ".c", ".h", ".hpp", ".cs", ".go", ".rs", ".php",
    ".rb", ".sql", ".sh", ".json", ".xml", ".yaml", ".yml", ".md", ".txt"
}

DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".txt"}

def extract_text_from_pdf(file_path: str) -> str:
    text_chunks = []
    try:
        reader = pypdf.PdfReader(file_path)
        for page_num, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(f"--- Page {page_num + 1} ---\n" + page_text.strip())
        return "\n\n".join(text_chunks)
    except Exception as e:
        # Fallback to PyMuPDF / fitz if available
        try:
            import fitz
            doc = fitz.open(file_path)
            for page in doc:
                text = page.get_text()
                if text:
                    text_chunks.append(text.strip())
            return "\n\n".join(text_chunks)
        except Exception:
            raise ValueError(f"Could not parse PDF document: {str(e)}")

def extract_text_from_docx(file_path: str) -> str:
    try:
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Could not parse Word DOCX document: {str(e)}")

def extract_text_from_code_or_text(file_path: str) -> str:
    encodings = ["utf-8", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
        except Exception as e:
            raise ValueError(f"Could not read source code/text file: {str(e)}")
    raise ValueError("File encoding could not be determined. Please ensure it is saved in UTF-8 format.")

def extract_text_from_zip(file_path: str, max_files: int = 40, max_file_size_kb: int = 500) -> Tuple[str, List[str]]:
    extracted_chunks = []
    manifest = []
    try:
        with zipfile.ZipFile(file_path, "r") as z:
            namelist = z.namelist()
            for name in namelist:
                # Security: prevent zip slip traversal
                if ".." in name or name.startswith("/") or name.startswith("\\"):
                    continue
                info = z.getinfo(name)
                if info.is_dir():
                    continue
                
                manifest.append(name)
                ext = Path(name).suffix.lower()
                if ext in CODE_EXTENSIONS or ext in DOCUMENT_EXTENSIONS:
                    if len(extracted_chunks) >= max_files:
                        extracted_chunks.append(f"\n[Note: Additional files truncated. Total files in zip: {len(namelist)}]")
                        break
                    
                    if info.file_size > max_file_size_kb * 1024:
                        extracted_chunks.append(f"\n--- File: {name} (Skipped: exceeds {max_file_size_kb}KB limit) ---")
                        continue
                    
                    try:
                        content_bytes = z.read(name)
                        try:
                            content = content_bytes.decode("utf-8")
                        except UnicodeDecodeError:
                            content = content_bytes.decode("latin-1", errors="replace")
                        extracted_chunks.append(f"\n================ File: {name} ================\n{content}")
                    except Exception as fe:
                        extracted_chunks.append(f"\n--- File: {name} (Error reading file: {str(fe)}) ---")
                        
        return "\n".join(extracted_chunks), manifest
    except zipfile.BadZipFile:
        raise ValueError("The uploaded ZIP archive is corrupted or not a valid ZIP file.")
    except Exception as e:
        raise ValueError(f"Failed to extract ZIP archive: {str(e)}")

def parse_and_verify_document(file_path: str, original_filename: str) -> Dict[str, Any]:
    """
    Parses document, extracts readable text, and validates against corruption/emptiness.
    """
    path = Path(file_path)
    if not path.exists():
        return {
            "success": False,
            "error_message": "Uploaded file not found on server storage.",
            "extracted_text": "",
            "file_type": "Unknown",
            "word_count": 0,
            "file_manifest": []
        }
    
    file_size = path.stat().st_size
    if file_size == 0:
        return {
            "success": False,
            "error_message": "Uploaded file is completely empty (0 bytes). Please upload a valid submission.",
            "extracted_text": "",
            "file_type": "Empty",
            "word_count": 0,
            "file_manifest": []
        }
    
    ext = Path(original_filename).suffix.lower()
    extracted_text = ""
    file_manifest = [original_filename]
    file_type = ext.replace(".", "").upper() or "UNKNOWN"
    
    try:
        if ext == ".pdf":
            file_type = "PDF"
            extracted_text = extract_text_from_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            file_type = "DOCX"
            extracted_text = extract_text_from_docx(file_path)
        elif ext == ".zip":
            file_type = "ZIP"
            extracted_text, file_manifest = extract_text_from_zip(file_path)
        elif ext in CODE_EXTENSIONS:
            file_type = f"Code/Text ({ext})"
            extracted_text = extract_text_from_code_or_text(file_path)
        else:
            return {
                "success": False,
                "error_message": f"Unsupported file extension '{ext}'. Permitted formats: PDF, DOCX, ZIP, or code files (.py, .js, .java, etc.).",
                "extracted_text": "",
                "file_type": file_type,
                "word_count": 0,
                "file_manifest": []
            }
    except Exception as e:
        return {
            "success": False,
            "error_message": f"Failed to extract text from document: {str(e)}",
            "extracted_text": "",
            "file_type": file_type,
            "word_count": 0,
            "file_manifest": file_manifest
        }
    
    cleaned_text = extracted_text.strip()
    words = cleaned_text.split()
    word_count = len(words)
    
    if word_count < 10 and file_type != "ZIP":
        return {
            "success": False,
            "error_message": f"The document contains insufficient readable text ({word_count} words). It might be an image-only scan or blank document.",
            "extracted_text": cleaned_text,
            "file_type": file_type,
            "word_count": word_count,
            "file_manifest": file_manifest
        }
        
    return {
        "success": True,
        "error_message": None,
        "extracted_text": cleaned_text,
        "file_type": file_type,
        "word_count": word_count,
        "file_manifest": file_manifest
    }
